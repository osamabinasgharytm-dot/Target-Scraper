"""Write combined product/review workbooks and optional review-range Excel files."""

from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.config import PRODUCT_COLUMNS, REVIEW_COLUMNS

logger = logging.getLogger("scrapper")


def _review_identity(row: dict) -> tuple[str, ...] | None:
    """Build a stable review key, with a content fallback for older exports."""
    dpci_value = row.get("DPCI")
    dpci = "" if dpci_value is None or pd.isna(dpci_value) else str(dpci_value).strip()
    review_id_value = row.get("Review ID")
    review_id = (
        ""
        if review_id_value is None or pd.isna(review_id_value)
        else str(review_id_value).strip()
    )
    if review_id:
        return ("id", dpci, review_id)

    def normalized(value: object) -> str:
        if value is None or pd.isna(value):
            return ""
        if isinstance(value, (date, datetime, pd.Timestamp)):
            return value.strftime("%Y-%m-%d")
        return " ".join(str(value).split()).casefold()

    date_value = row.get("Date Posted")
    parsed_date = pd.to_datetime(date_value, dayfirst=True, errors="coerce")
    signature = (
        normalized(row.get("DPCI")),
        "" if pd.isna(parsed_date) else parsed_date.strftime("%Y-%m-%d"),
        normalized(row.get("Review Rating")),
        normalized(row.get("Review Title")),
        normalized(row.get("Review Text")),
    )
    if not any(signature[1:]):
        return None
    return ("content", *signature)


def load_website_error_dpcis(path: Path) -> list[str]:
    """Find DPCIs whose latest Products-sheet entry has Website Status Error."""
    if not path.is_file():
        raise FileNotFoundError(f"Excel workbook not found: {path}")

    with pd.ExcelFile(path) as workbook:
        if "Products" not in workbook.sheet_names:
            raise ValueError(f"Excel workbook has no Products sheet: {path}")
        products = pd.read_excel(
            workbook,
            sheet_name="Products",
            dtype={"DPCI": str},
        )

    missing_columns = {"DPCI", "Website Status"}.difference(products.columns)
    if missing_columns:
        missing = ", ".join(sorted(missing_columns))
        raise ValueError(
            f"Products sheet must contain DPCI and Website Status columns; "
            f"missing: {missing}"
        )

    latest_status: dict[str, str] = {}
    for dpci, status in zip(products["DPCI"], products["Website Status"]):
        if pd.isna(dpci) or pd.isna(status):
            continue
        dpci_value = str(dpci).strip()
        if dpci_value:
            latest_status[dpci_value] = str(status).strip()

    return [
        dpci
        for dpci, status in latest_status.items()
        if status.casefold() == "error"
    ]


def _ensure_headers(sheet, expected_columns: list[str]) -> dict[str, int]:
    headers = {
        str(sheet.cell(row=1, column=column).value).strip(): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(row=1, column=column).value is not None
    }
    for column in expected_columns:
        if column not in headers:
            new_column = sheet.max_column + 1
            sheet.cell(row=1, column=new_column, value=column)
            headers[column] = new_column
    return headers


def append_products_and_reviews_excel(
    product_rows: list[dict],
    review_rows: list[dict],
    path: Path,
) -> Path:
    """Update unique product rows and append only reviews missing from the workbook."""
    if not path.is_file():
        raise FileNotFoundError(f"Excel workbook not found: {path}")

    workbook = load_workbook(path)
    required_sheets = {"Products", "Reviews"}
    missing_sheets = required_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        missing = ", ".join(sorted(missing_sheets))
        raise ValueError(
            f"Excel workbook must contain Products and Reviews sheets; "
            f"missing: {missing}"
        )

    appended_review_count = 0
    product_sheet = workbook["Products"]
    product_headers = _ensure_headers(product_sheet, PRODUCT_COLUMNS)
    if "DPCI" not in product_headers or "Website Status" not in product_headers:
        raise ValueError("Products sheet must contain DPCI and Website Status columns")

    def write_row(sheet, headers: dict[str, int], row_number: int, row: dict) -> None:
        for column, value in row.items():
            if column not in headers:
                continue
            if value is None or pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            sheet.cell(row=row_number, column=headers[column], value=value)

    product_rows_by_dpci: dict[str, list[int]] = {}
    for row_number in range(2, product_sheet.max_row + 1):
        dpci = product_sheet.cell(
            row=row_number,
            column=product_headers["DPCI"],
        ).value
        if dpci is None or not str(dpci).strip():
            continue
        dpci_key = str(dpci).strip()
        product_rows_by_dpci.setdefault(dpci_key, []).append(row_number)

    duplicate_product_rows = [
        row_number
        for row_numbers in product_rows_by_dpci.values()
        for row_number in row_numbers[:-1]
    ]
    for row_number in sorted(duplicate_product_rows, reverse=True):
        product_sheet.delete_rows(row_number)

    product_row_numbers: dict[str, int] = {}
    for row_number in range(2, product_sheet.max_row + 1):
        dpci = product_sheet.cell(
            row=row_number,
            column=product_headers["DPCI"],
        ).value
        if dpci is not None and str(dpci).strip():
            product_row_numbers[str(dpci).strip()] = row_number

    for row in product_rows:
        dpci = str(row.get("DPCI") or "").strip()
        if not dpci:
            product_sheet.append([None] * product_sheet.max_column)
            write_row(product_sheet, product_headers, product_sheet.max_row, row)
            continue
        if dpci in product_row_numbers:
            write_row(product_sheet, product_headers, product_row_numbers[dpci], row)
        else:
            product_sheet.append([None] * product_sheet.max_column)
            row_number = product_sheet.max_row
            write_row(product_sheet, product_headers, row_number, row)
            product_row_numbers[dpci] = row_number

    review_sheet = workbook["Reviews"]
    review_headers = _ensure_headers(review_sheet, REVIEW_COLUMNS)
    seen_review_ids: set[tuple[str, ...]] = set()
    seen_review_content: set[tuple[str, ...]] = set()
    duplicate_review_rows: list[int] = []
    for row_number in range(2, review_sheet.max_row + 1):
        existing_row = {
            header: review_sheet.cell(row=row_number, column=column).value
            for header, column in review_headers.items()
        }
        identity = _review_identity(existing_row)
        if identity is None:
            continue
        content_identity = _review_identity(
            {
                **existing_row,
                "Review ID": "",
            }
        )
        already_exists = (
            identity in seen_review_ids
            or (
                content_identity is not None
                and content_identity in seen_review_content
            )
        )
        if already_exists:
            duplicate_review_rows.append(row_number)
            continue
        if content_identity is not None:
            seen_review_content.add(content_identity)
        if identity[0] == "id":
            seen_review_ids.add(identity)

    for row_number in reversed(duplicate_review_rows):
        review_sheet.delete_rows(row_number)

    for row in review_rows:
        identity = _review_identity(row)
        if identity is not None:
            content_identity = _review_identity({**row, "Review ID": ""})
            if (
                identity in seen_review_ids
                or (
                    content_identity is not None
                    and content_identity in seen_review_content
                )
            ):
                continue
        review_sheet.append([None] * review_sheet.max_column)
        write_row(review_sheet, review_headers, review_sheet.max_row, row)
        if identity is not None:
            if content_identity is not None:
                seen_review_content.add(content_identity)
            if identity[0] == "id":
                seen_review_ids.add(identity)
        appended_review_count += 1

    workbook.save(path)
    logger.info(
        "Updated %d product rows and added %d new review rows to %s",
        len(product_rows),
        appended_review_count,
        path,
    )
    return path


# def write_products_csv(rows: list[dict], path: Path) -> Path:
    # path.parent.mkdir(parents=True, exist_ok=True)
    # df = pd.DataFrame(rows, columns=PRODUCT_COLUMNS)
    # df.to_csv(path, index=False, encoding="utf-8-sig")
    # logger.info("Wrote %d product rows -> %s", len(df), path)
    # return path


# def write_reviews_csv(rows: list[dict], path: Path) -> Path:
    # path.parent.mkdir(parents=True, exist_ok=True)
    # df = pd.DataFrame(rows, columns=REVIEW_COLUMNS)
    # df.to_csv(path, index=False, encoding="utf-8-sig")
    # logger.info("Wrote %d review rows -> %s", len(df), path)
    # return path


def write_reviews_excel(rows: list[dict], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows, columns=REVIEW_COLUMNS)
    df.to_excel(path, index=False)
    logger.info("Wrote %d review rows -> %s", len(df), path)
    return path


def write_products_and_reviews_excel(
    product_rows: list[dict],
    review_rows: list[dict],
    path: Path,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    products_df = pd.DataFrame(product_rows, columns=PRODUCT_COLUMNS)
    reviews_df = pd.DataFrame(review_rows, columns=REVIEW_COLUMNS)
    with pd.ExcelWriter(path) as writer:
        products_df.to_excel(writer, sheet_name="Products", index=False)
        reviews_df.to_excel(writer, sheet_name="Reviews", index=False)
    logger.info(
        "Wrote %d product rows and %d review rows -> %s",
        len(products_df),
        len(reviews_df),
        path,
    )
    return path
